"""Backend additions for the web app: thumbnails, template download, demo server."""

import io
import tempfile
from pathlib import Path

import pytest
from openpyxl import load_workbook
from PIL import Image

from backend.app import demo
from backend.app.queries import load_script
from backend.grading.scheme import load_scheme
from tests.app.conftest import add_user, login, make_client


async def test_scheme_template_download(app, client):
    await add_user(app, "ADM1", role="admin")
    await login(client, "ADM1")
    response = await client.get("/api/admin/scheme-template")
    assert response.status_code == 200
    assert "marking_scheme_template.xlsx" in response.headers["content-disposition"]
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "t.xlsx"
        path.write_bytes(response.content)
        assert len(load_scheme(path).items) == 5  # the template's example rows load cleanly


async def test_scheme_template_admin_only(app, client):
    await add_user(app, "T100")
    await login(client, "T100")
    assert (await client.get("/api/admin/scheme-template")).status_code == 403


def test_demo_refuses_without_confirmation(capsys):
    assert demo.main([]) == 2
    assert "--yes-demo" in capsys.readouterr().err


def test_demo_samples(tmp_path):
    demo.write_samples(tmp_path)
    scheme = load_scheme(tmp_path / "sample_scheme.xlsx")
    assert [i.key for i in scheme.items] == ["1", "2", "3"]
    assert (tmp_path / "sample_script.pdf").read_bytes().startswith(b"%PDF-")


@pytest.fixture
async def demo_app(tmp_path):
    settings = demo.demo_settings(tmp_path)
    await demo.seed(settings)
    application = demo.build_demo_app(settings)
    async with application.router.lifespan_context(application):
        yield application


async def test_demo_server_grades_with_fake_ai_and_makes_thumbnails(demo_app, tmp_path):
    demo.write_samples(tmp_path)
    async with make_client(demo_app) as admin, make_client(demo_app) as teacher:
        await login(admin, "ADM1", "demo-admin-1")
        exam = (await admin.post(
            "/api/admin/exams", data={"name": "Demo", "subject": "Physics", "class_section": "A"},
            files={"scheme": ("s.xlsx", (tmp_path / "sample_scheme.xlsx").read_bytes())},
        )).json()
        bundle = (await admin.post(f"/api/admin/exams/{exam['id']}/bundles",
                                   json={"code": "B1", "teacher_employee_id": "T100"})).json()
        await admin.post(f"/api/admin/bundles/{bundle['id']}/scripts", files=[
            ("files", ("s.pdf", (tmp_path / "sample_script.pdf").read_bytes(), "application/pdf"))])
        await demo_app.state.runner.join()

        async with demo_app.state.sessionmaker() as db:
            script = await load_script(db, 1)
        assert script.status == "graded" and script.roll_number is None
        marks = {q.question: (q.ai_marks, q.needs_review) for q in script.questions}
        assert marks == {"1": (1, False), "2": (2, False), "3": (2, True)}

        await login(teacher, "T100", "demo-teacher-1")
        full = await teacher.get(f"/api/my/scripts/{script.id}/pages/2")
        thumb = await teacher.get(f"/api/my/scripts/{script.id}/pages/2?size=thumb")
        assert full.status_code == thumb.status_code == 200
        assert max(Image.open(io.BytesIO(thumb.content)).size) == 240
        assert max(Image.open(io.BytesIO(full.content)).size) > 240
        bad = await teacher.get(f"/api/my/scripts/{script.id}/pages/2?size=huge")
        assert bad.status_code == 422


async def test_thumbnail_falls_back_to_full_page(demo_app, tmp_path):
    """Scripts graded before thumbnails existed still show something in the strip."""
    demo.write_samples(tmp_path)
    async with make_client(demo_app) as admin, make_client(demo_app) as teacher:
        await login(admin, "ADM1", "demo-admin-1")
        exam = (await admin.post(
            "/api/admin/exams", data={"name": "Demo", "subject": "P", "class_section": "A"},
            files={"scheme": ("s.xlsx", (tmp_path / "sample_scheme.xlsx").read_bytes())},
        )).json()
        bundle = (await admin.post(f"/api/admin/exams/{exam['id']}/bundles",
                                   json={"code": "B1", "teacher_employee_id": "T100"})).json()
        await admin.post(f"/api/admin/bundles/{bundle['id']}/scripts", files=[
            ("files", ("s.pdf", (tmp_path / "sample_script.pdf").read_bytes(), "application/pdf"))])
        await demo_app.state.runner.join()
        thumb_path = (demo_app.state.settings.storage_dir / "scripts" / "1" / "pages"
                      / "thumb_1.jpg")
        thumb_path.unlink()
        await login(teacher, "T100", "demo-teacher-1")
        response = await teacher.get("/api/my/scripts/1/pages/1?size=thumb")
        assert response.status_code == 200
        assert max(Image.open(io.BytesIO(response.content)).size) > 240


def test_demo_never_uses_env_database(monkeypatch, tmp_path):
    monkeypatch.setenv("DATABASE_URL", "postgresql://real:secret@supabase/postgres")
    settings = demo.demo_settings(tmp_path)
    assert settings.database_url.startswith("sqlite")


def test_template_workbook_has_instructions(tmp_path):
    from backend.grading.template import write_template

    write_template(tmp_path / "t.xlsx")
    assert "Instructions" in load_workbook(tmp_path / "t.xlsx").sheetnames
