"""End-to-end over the API: upload -> AI grading -> teacher review -> submit -> Excel."""

from datetime import UTC, datetime

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from backend.app.excel_export import ExportError, export_exam
from backend.app.marks import approval_problems, counted_ids, total
from backend.app.models import AuditLog, Bundle, Exam, QuestionMark, Script
from backend.app.queries import load_exam_full
from backend.extract.vision import VisionClient
from backend.grading.grader import Grader
from tests.app.conftest import add_user, app_settings, login, make_client
from tests.app.factories import create_exam, script_pdf
from tests.fakes import FakeGroq, message, page_number_of


def vision_handler(request):
    assert page_number_of(request) == 2
    return message({"segments": [
        {"question_number": "1", "text": "(c)", "has_diagram": False, "illegible": False,
         "illegible_notes": []},
        {"question_number": "5", "text": "two of the points", "has_diagram": False,
         "illegible": False, "illegible_notes": []}]})


def grade_handler(request):
    # A deduction with medium confidence: flagged for the teacher to check.
    return message({"marks_awarded": 2, "points_met": ["a", "b"], "points_missing": ["c"],
                    "reason": "Two of three points.", "confidence": "medium",
                    "ocr_problem": False})


@pytest.fixture
def settings(tmp_path):
    return app_settings(tmp_path, transcribe_provider="groq", page_rotation="0",
                        vision_backoff_base_s=0, vision_backoff_max_s=0, vision_max_retries=0)


@pytest.fixture
def app_factories():
    return {"extraction_factory": lambda s: VisionClient(FakeGroq(vision_handler), s),
            "grader_factory": lambda s: Grader(FakeGroq(grade_handler), s)}


async def setup_exam(app, client, bundles=(("B1", "T100", ["21CS045"]),)):
    """Admin creates an exam, bundles and uploads scripts; waits for AI grading."""
    await add_user(app, "ADM1", role="admin")
    for eid in {b[1] for b in bundles} | {"T200"}:
        await add_user(app, eid)
    await login(client, "ADM1")
    exam = (await create_exam(client)).json()
    ids = {}
    for code, teacher, rolls in bundles:
        bundle = (await client.post(f"/api/admin/exams/{exam['id']}/bundles",
                                    json={"code": code, "teacher_employee_id": teacher})).json()
        files = [("files", (f"{r}.pdf", script_pdf(r), "application/pdf")) for r in rolls]
        await client.post(f"/api/admin/bundles/{bundle['id']}/scripts", files=files)
        ids[code] = bundle["id"]
    await app.state.runner.join()
    await client.post("/api/auth/logout")
    return exam, ids


async def first_script(teacher, bundle_id):
    scripts = (await teacher.get(f"/api/my/bundles/{bundle_id}")).json()["scripts"]
    return (await teacher.get(f"/api/my/scripts/{scripts[0]['id']}")).json()


def q_of(script, label):
    return next(q for q in script["questions"] if q["question"] == label)


async def review_and_approve(teacher, script, q5_marks=2.5):
    q5 = q_of(script, "5")
    r = await teacher.patch(f"/api/my/scripts/{script['id']}/questions/{q5['id']}",
                            json={"final_marks": q5_marks, "teacher_comment": "partly right"})
    assert r.status_code == 200, r.text
    r = await teacher.post(f"/api/my/scripts/{script['id']}/approve")
    assert r.status_code == 200, r.text
    return r.json()


# --- access ----------------------------------------------------------------------------


async def test_teachers_see_only_their_bundles(app, client):
    _, ids = await setup_exam(app, client)
    async with make_client(app) as asha, make_client(app) as ravi:
        await login(asha, "T100")
        await login(ravi, "T200")
        mine = (await asha.get("/api/my/bundles")).json()
        assert [b["code"] for b in mine] == ["B1"]
        assert mine[0]["graded"] == 1 and mine[0]["can_submit"] is False
        assert (await ravi.get("/api/my/bundles")).json() == []
        script = await first_script(asha, ids["B1"])
        assert (await ravi.get(f"/api/my/bundles/{ids['B1']}")).status_code == 404
        assert (await ravi.get(f"/api/my/scripts/{script['id']}")).status_code == 404
        assert (await ravi.get(f"/api/my/scripts/{script['id']}/pages/1")).status_code == 404
        assert (await ravi.post(f"/api/my/scripts/{script['id']}/approve")).status_code == 404


async def test_admin_cannot_use_teacher_review(app, client):
    await setup_exam(app, client)
    await login(client, "ADM1")
    assert (await client.get("/api/my/bundles")).status_code == 403


# --- reviewing a script -----------------------------------------------------------------


async def test_script_view_has_marks_pages_and_flags(app, client):
    _, ids = await setup_exam(app, client)
    async with make_client(app) as asha:
        await login(asha, "T100")
        s = await first_script(asha, ids["B1"])
        assert (s["roll_number"], s["page_count"], s["editable"]) == ("21CS045", 2, True)
        assert (s["ai_total"], s["final_total"], s["max_marks"]) == (3, 3, 4)
        q5 = q_of(s, "5")
        assert q5["needs_review"] and not q5["confirmed"]
        assert "marks deducted with less than high confidence" in q5["review_notes"]
        assert s["problems"] == ["Check question 5 (flagged for review)."]
        page = await asha.get(f"/api/my/scripts/{s['id']}/pages/1")
        assert page.status_code == 200 and page.content[:2] == b"\xff\xd8"
        assert page.headers["cache-control"] == "private, max-age=300"
        assert (await asha.get(f"/api/my/scripts/{s['id']}/pages/9")).status_code == 404


@pytest.mark.parametrize("bad", [3.5, -1, 1.25])
async def test_invalid_marks_rejected(app, client, bad):
    _, ids = await setup_exam(app, client)
    async with make_client(app) as asha:
        await login(asha, "T100")
        s = await first_script(asha, ids["B1"])
        r = await asha.patch(f"/api/my/scripts/{s['id']}/questions/{q_of(s, '5')['id']}",
                             json={"final_marks": bad})
        assert r.status_code == 422
        assert "between 0 and 3" in r.json()["detail"]


async def test_edit_mark_is_audited_and_confirms(app, client):
    _, ids = await setup_exam(app, client)
    async with make_client(app) as asha:
        await login(asha, "T100")
        s = await first_script(asha, ids["B1"])
        q5 = q_of(s, "5")
        r = await asha.patch(f"/api/my/scripts/{s['id']}/questions/{q5['id']}",
                             json={"final_marks": 2.5, "teacher_comment": "partly right"})
        updated = r.json()
        new_q5 = q_of(updated, "5")
        assert (new_q5["final_marks"], new_q5["ai_marks"]) == (2.5, 2)
        assert new_q5["confirmed"] and new_q5["changed"]
        assert new_q5["teacher_comment"] == "partly right"
        assert updated["final_total"] == 3.5 and updated["problems"] == []
    async with app.state.sessionmaker() as db:
        entry = await db.scalar(select(AuditLog).where(AuditLog.action == "edit_mark"))
    assert entry.detail["old"] == {"final_marks": 2, "confirmed": False}
    assert entry.detail["new"] == {"final_marks": 2.5, "confirmed": True}


async def test_approval_rules(app, client):
    _, ids = await setup_exam(app, client)
    async with make_client(app) as asha:
        await login(asha, "T100")
        s = await first_script(asha, ids["B1"])
        blocked = await asha.post(f"/api/my/scripts/{s['id']}/approve")
        assert blocked.status_code == 409
        assert blocked.json()["detail"]["problems"] == ["Check question 5 (flagged for review)."]
        # Confirming without changing the mark is enough.
        await asha.patch(f"/api/my/scripts/{s['id']}/questions/{q_of(s, '5')['id']}",
                         json={"confirmed": True})
        approved = (await asha.post(f"/api/my/scripts/{s['id']}/approve")).json()
        assert (approved["status"], approved["editable"]) == ("approved", False)
        locked = await asha.patch(f"/api/my/scripts/{s['id']}/questions/{q_of(s, '5')['id']}",
                                  json={"final_marks": 1})
        assert locked.status_code == 409
        reopened = (await asha.post(f"/api/my/scripts/{s['id']}/reopen")).json()
        assert reopened["status"] == "graded"


async def test_roll_number_required_and_editable(app, client):
    _, ids = await setup_exam(app, client)
    async with app.state.sessionmaker() as db:
        script = await db.scalar(select(Script))
        script.roll_number = None
        await db.commit()
    async with make_client(app) as asha:
        await login(asha, "T100")
        s = await first_script(asha, ids["B1"])
        assert "Enter the roll number." in s["problems"]
        bad = await asha.patch(f"/api/my/scripts/{s['id']}", json={"roll_number": "21 CS!"})
        assert bad.status_code == 422
        fixed = (await asha.patch(f"/api/my/scripts/{s['id']}",
                                  json={"roll_number": "21cs099"})).json()
        assert (fixed["roll_number"], fixed["roll_number_source"]) == ("21CS099", "teacher")


# --- submit and export -----------------------------------------------------------------


async def test_submit_last_bundle_writes_excel(app, client, settings):
    exam, ids = await setup_exam(app, client)
    async with make_client(app) as asha:
        await login(asha, "T100")
        early = await asha.post(f"/api/my/bundles/{ids['B1']}/submit")
        assert early.status_code == 409
        await review_and_approve(asha, await first_script(asha, ids["B1"]))
        result = (await asha.post(f"/api/my/bundles/{ids['B1']}/submit")).json()
        assert result["bundle"]["status"] == "submitted"
        assert result["exam_completed"] is True and result["exported"] is True
        s = await first_script(asha, ids["B1"])
        r = await asha.patch(f"/api/my/scripts/{s['id']}/questions/{q_of(s, '5')['id']}",
                             json={"final_marks": 1})
        assert r.status_code == 409
        assert (await asha.post(f"/api/my/scripts/{s['id']}/reopen")).status_code == 409

    files = list(settings.export_dir.glob("*.xlsx"))
    assert len(files) == 1
    ws = load_workbook(files[0]).active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    assert rows[0] == ["Roll No", "1", "5", "Total"]
    assert rows[1] == ["Max marks", 1, 3, 4]
    assert rows[2] == ["21CS045", 1, 2.5, 3.5]
    assert ws.title == "Physics XII-A Mid 1"

    await login(client, "ADM1")
    detail = (await client.get(f"/api/admin/exams/{exam['id']}")).json()
    assert detail["status"] == "exported" and detail["export_error"] is None
    download = await client.get(f"/api/admin/exams/{exam['id']}/export/file")
    assert download.status_code == 200 and download.content[:2] == b"PK"


async def test_no_export_until_every_bundle_is_submitted(app, client, settings):
    exam, ids = await setup_exam(app, client, bundles=(("B1", "T100", ["21CS045"]),
                                                       ("B2", "T200", ["21CS046"])))
    async with make_client(app) as asha:
        await login(asha, "T100")
        await review_and_approve(asha, await first_script(asha, ids["B1"]))
        result = (await asha.post(f"/api/my/bundles/{ids['B1']}/submit")).json()
        assert (result["exam_completed"], result["exported"]) == (False, False)
    assert list(settings.export_dir.glob("*.xlsx")) == []
    await login(client, "ADM1")
    forced = await client.post(f"/api/admin/exams/{exam['id']}/export")
    assert forced.status_code == 409
    assert "Bundle B2 is not submitted yet." in forced.json()["detail"]["problems"]
    assert list(settings.export_dir.glob("*.xlsx")) == []


async def test_duplicate_roll_numbers_block_export(app, client, settings):
    exam, ids = await setup_exam(app, client, bundles=(("B1", "T100", ["21CS045"]),
                                                       ("B2", "T200", ["21CS045"])))
    for teacher, code in (("T100", "B1"), ("T200", "B2")):
        async with make_client(app) as t:
            await login(t, teacher)
            await review_and_approve(t, await first_script(t, ids[code]))
            result = (await t.post(f"/api/my/bundles/{ids[code]}/submit")).json()
    assert result["exam_completed"] is True and result["exported"] is False
    assert list(settings.export_dir.glob("*.xlsx")) == []
    await login(client, "ADM1")
    detail = (await client.get(f"/api/admin/exams/{exam['id']}")).json()
    assert "Roll number 21CS045 appears twice (bundles B1 and B2)." in detail["export_error"]


async def test_master_workbook_keeps_other_sheets(app, client, settings, tmp_path):
    master = tmp_path / "college_marks.xlsx"
    wb = Workbook()
    wb.active.title = "Chemistry XII-A Mid 1"
    wb.active["A1"] = "keep me"
    wb.save(master)
    settings.excel_master_path = master
    _, ids = await setup_exam(app, client)
    async with make_client(app) as asha:
        await login(asha, "T100")
        await review_and_approve(asha, await first_script(asha, ids["B1"]))
        await asha.post(f"/api/my/bundles/{ids['B1']}/submit")
    book = load_workbook(master)
    assert book.sheetnames == ["Chemistry XII-A Mid 1", "Physics XII-A Mid 1"]
    assert book["Chemistry XII-A Mid 1"]["A1"].value == "keep me"
    assert book["Physics XII-A Mid 1"]["D3"].value == 3.5


# --- the hard rule, at the lowest level ---------------------------------------------------


async def test_export_refuses_unapproved_marks(app, client, settings):
    exam, ids = await setup_exam(app, client)
    async with app.state.sessionmaker() as db:
        bundle = await db.get(Bundle, ids["B1"])
        bundle.submitted_at = datetime.now(UTC)  # submitted, but the script is not approved
        await db.commit()
        full = await load_exam_full(db, exam["id"])
    with pytest.raises(ExportError) as exc:
        export_exam(full, export_dir=settings.export_dir, master_path=None)
    assert exc.value.problems == ["A script in bundle B1 is not approved."]


# --- OR groups and totals ---------------------------------------------------------------


def qm(id_, final, group=None, ai=None, needs_review=False, confirmed=False, label=None):
    return QuestionMark(id=id_, question=label or str(id_), final_marks=final, ai_marks=ai,
                        choice_group=group, needs_review=needs_review, confirmed=confirmed,
                        position=id_, key=str(id_), qtype="long", max_marks=5, ai_status="graded")


def test_or_group_counts_the_better_alternative():
    qs = [qm(1, 2), qm(2, 3, "g"), qm(3, 4, "g")]
    assert counted_ids(qs) == {1, 3}
    assert total(qs) == 6


def test_approval_ignores_unused_or_alternative():
    qs = [qm(1, 2), qm(2, None, "g", needs_review=True), qm(3, 4, "g")]
    assert approval_problems(qs, "21CS045") == []
    qs[0].final_marks = None
    assert approval_problems(qs, None) == ["Enter the roll number.",
                                           "Enter marks for question 1."]


async def test_exam_model_relations_load(app):
    async with app.state.sessionmaker() as db:
        assert (await db.scalars(select(Exam))).all() == []
