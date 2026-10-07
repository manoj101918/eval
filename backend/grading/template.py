"""Create a blank marking-scheme Excel template.

    python -m backend.grading.template scheme.xlsx
"""

import argparse
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from backend.grading.scheme import QUESTION_TYPES

COLUMNS = [
    ("Question", 12, "As written in the question paper: 5, 12(b), 34(a)(i)"),
    ("Max marks", 10, "Number; halves allowed (e.g. 1.5)"),
    ("Type", 11, "mcq / short / long / numerical"),
    ("Correct option", 10, "MCQ only: a, b, c, d or e"),
    ("Model answer / key points", 60, "The expected answer. List the points that earn marks."),
    ("Marking guidance", 45, "How to split marks, e.g. 1 mark formula, 1 substitution, 1 answer"),
    ("Choice group", 12, "Same text for OR alternatives; only one is counted"),
]
EXAMPLES = [
    ["1", 1, "mcq", "c", "", "", ""],
    ["5(a)", 2, "short", "",
     "Ohm's law: current through a conductor is proportional to the potential difference "
     "across it at constant temperature (V = IR).",
     "1 mark statement, 1 mark formula", ""],
    ["12", 3, "numerical", "",
     "Reff = R1 + R2 = 6 ohm; I = V / Reff = 6 / 6 = 1 A",
     "1 mark formula, 1 mark substitution, 1 mark answer with unit", ""],
    ["31", 5, "long", "", "Interference vs diffraction: (1) intensity of bright fringes ...",
     "1 mark per correct difference, max 5", "31-OR"],
    ["31 OR", 5, "long", "", "Alternative question 31: ...", "", "31-OR"],
]
INSTRUCTIONS = [
    "How to fill in this marking scheme",
    "",
    "1. One row per question or sub-part that gets its own marks (e.g. 34(a)(i)).",
    "2. Question numbers must match the question paper, so they can be matched to the "
    "student's answers.",
    "3. Type: mcq (graded automatically from 'Correct option'), short, long or numerical.",
    "4. For non-MCQ rows, write the model answer or the key points, and how marks are split.",
    "   The AI follows the scheme and gives partial credit per key point; spelling and grammar "
    "are not penalised unless the guidance says so.",
    "5. 'Attempt any one' (OR) questions: write the alternative as e.g. '31 OR', give both "
    "the same Choice group text and the same max marks. The one the student answered counts.",
    "6. Delete the example rows before use. Blank rows are ignored.",
    "",
    "AI marks are proposals only: a teacher reviews every mark before it is final.",
]


def write_template(path: Path) -> None:
    wb = Workbook()
    scheme = wb.active
    scheme.title = "Scheme"
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    for col, (name, width, note) in enumerate(COLUMNS, start=1):
        cell = scheme.cell(row=1, column=col, value=name)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        scheme.column_dimensions[cell.column_letter].width = width
        cell.comment = Comment(note, "Answer Script Evaluator")  # hover hint, not a data row
    scheme.freeze_panes = "A2"
    for r, values in enumerate(EXAMPLES, start=2):
        for c, value in enumerate(values, start=1):
            scheme.cell(row=r, column=c, value=value).alignment = Alignment(
                wrap_text=True, vertical="top")

    types = DataValidation(type="list", formula1=f'"{",".join(QUESTION_TYPES)}"',
                           allow_blank=True)
    scheme.add_data_validation(types)
    types.add("C2:C500")
    options = DataValidation(type="list", formula1='"a,b,c,d,e"', allow_blank=True)
    scheme.add_data_validation(options)
    options.add("D2:D500")

    exam = wb.create_sheet("Exam")
    for r, (name, value) in enumerate([("Subject", ""), ("Exam", "")], start=1):
        exam.cell(row=r, column=1, value=name).font = Font(bold=True)
        exam.cell(row=r, column=2, value=value)
    exam.column_dimensions["A"].width = 12
    exam.column_dimensions["B"].width = 40

    info = wb.create_sheet("Instructions")
    for r, line in enumerate(INSTRUCTIONS, start=1):
        info.cell(row=r, column=1, value=line).font = Font(bold=(r == 1))
    info.column_dimensions["A"].width = 110
    wb.save(path)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="python -m backend.grading.template")
    parser.add_argument("path", type=Path, help="where to write the .xlsx template")
    args = parser.parse_args(argv)
    if args.path.exists():
        print(f"error: {args.path} already exists", file=sys.stderr)
        return 1
    write_template(args.path)
    print(f"Template written to {args.path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
