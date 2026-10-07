"""Marking scheme: read and validate the teacher's Excel file.

Sheet "Scheme" (or the first sheet) has a header row and one row per question or
sub-part. Columns are matched by header name, so their order does not matter:

  Question | Max marks | Type | Correct option | Model answer / key points |
  Marking guidance | Choice group

An optional "Exam" sheet holds "Subject" / "Exam" in its first two columns.
"""

import re
from pathlib import Path
from typing import Literal

from openpyxl import load_workbook
from pydantic import BaseModel, Field, model_validator

from backend.extract.normalize import normalize_question_label

QuestionType = Literal["mcq", "short", "long", "numerical"]
QUESTION_TYPES: tuple[str, ...] = ("mcq", "short", "long", "numerical")
MCQ_OPTIONS = "abcde"
HEADER_SEARCH_ROWS = 10

# Header synonyms -> field name (compared lower-case, punctuation removed)
HEADERS = {
    "question": "label", "questionno": "label", "qno": "label", "q": "label",
    "maxmarks": "max_marks", "marks": "max_marks", "maximummarks": "max_marks",
    "type": "qtype", "questiontype": "qtype",
    "correctoption": "correct_option", "option": "correct_option", "answeroption": "correct_option",
    "correctanswer": "correct_option",
    "modelanswerkeypoints": "model_answer", "modelanswer": "model_answer",
    "keypoints": "model_answer", "answer": "model_answer",
    "markingguidance": "guidance", "guidance": "guidance", "notes": "guidance",
    "choicegroup": "choice_group", "orgroup": "choice_group", "choice": "choice_group",
}
REQUIRED = ("label", "max_marks", "qtype")


class SchemeError(Exception):
    """The marking scheme file is invalid. `problems` lists each issue with its row."""

    def __init__(self, problems: list[str]) -> None:
        super().__init__("Marking scheme has problems:\n- " + "\n- ".join(problems))
        self.problems = problems


class SchemeItem(BaseModel):
    row: int
    label: str  # as the teacher wrote it, e.g. "34(a)(i)"
    key: str  # normalised, e.g. "34ai" (same normaliser as extraction)
    max_marks: float = Field(gt=0)
    qtype: QuestionType
    correct_option: str | None = None
    model_answer: str = ""
    guidance: str = ""
    choice_group: str | None = None

    @property
    def student_key(self) -> str:
        """Key the student writes: an OR alternative "31 OR" is answered as "31"."""
        if self.choice_group and self.key.endswith("or") and len(self.key) > 2:
            return self.key[:-2]
        return self.key

    @model_validator(mode="after")
    def _check(self) -> "SchemeItem":
        if (self.max_marks * 2) % 1:
            raise ValueError("max marks must be a multiple of 0.5")
        if self.qtype == "mcq":
            if not self.correct_option or self.correct_option not in MCQ_OPTIONS:
                raise ValueError("MCQ needs a correct option a-e")
        elif not self.model_answer.strip():
            raise ValueError("model answer / key points are required")
        return self


class MarkingScheme(BaseModel):
    subject: str | None = None
    exam: str | None = None
    items: list[SchemeItem]

    @property
    def total_marks(self) -> float:
        """Each choice group (OR alternatives) counts once."""
        seen: set[str] = set()
        total = 0.0
        for item in self.items:
            if item.choice_group:
                if item.choice_group in seen:
                    continue
                seen.add(item.choice_group)
            total += item.max_marks
        return total

    def by_key(self) -> dict[str, SchemeItem]:
        return {item.key: item for item in self.items}


def _header_key(value: object) -> str | None:
    if value is None:
        return None
    return HEADERS.get(re.sub(r"[^a-z]", "", str(value).lower()))


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _option(value: str) -> str | None:
    letters = re.sub(r"[^a-z]", "", value.lower())
    return letters or None


def load_scheme(path: Path) -> MarkingScheme:
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several types for bad files
        raise SchemeError([f"cannot open the file as Excel (.xlsx): {type(exc).__name__}"]) \
            from None
    try:
        sheet = wb["Scheme"] if "Scheme" in wb.sheetnames else wb.worksheets[0]
        rows = list(sheet.iter_rows(values_only=True))
        subject = exam = None
        if "Exam" in wb.sheetnames:
            for r in wb["Exam"].iter_rows(values_only=True):
                name = _cell_text(r[0]).lower() if r else ""
                value = _cell_text(r[1]) if r and len(r) > 1 else ""
                if name == "subject":
                    subject = value or None
                elif name == "exam":
                    exam = value or None
    finally:
        wb.close()

    header_row, columns = None, {}
    for i, r in enumerate(rows[:HEADER_SEARCH_ROWS]):
        found = {_header_key(v): j for j, v in enumerate(r) if _header_key(v)}
        if "label" in found and "max_marks" in found:
            header_row, columns = i, found
            break
    if header_row is None:
        raise SchemeError(["no header row with 'Question' and 'Max marks' columns found"])
    missing = [name for name in REQUIRED if name not in columns]
    if missing:
        raise SchemeError([f"missing column(s): {', '.join(missing)}"])

    problems: list[str] = []
    items: list[SchemeItem] = []
    for i, r in enumerate(rows[header_row + 1:], start=header_row + 2):
        def get(field: str, row: tuple = r) -> str:
            j = columns.get(field)
            return _cell_text(row[j]) if j is not None and j < len(row) else ""

        label = get("label")
        if not label and not any(_cell_text(v) for v in r):
            continue  # blank row
        if not label:
            problems.append(f"row {i}: Question is empty")
            continue
        key = normalize_question_label(label)
        if key is None:
            problems.append(f"row {i}: cannot read question label '{label}'")
            continue
        try:
            max_marks = float(get("max_marks"))
        except ValueError:
            problems.append(f"row {i} ({label}): Max marks must be a number")
            continue
        qtype = get("qtype").lower()
        if qtype not in QUESTION_TYPES:
            problems.append(f"row {i} ({label}): Type must be one of {', '.join(QUESTION_TYPES)}")
            continue
        try:
            items.append(SchemeItem(
                row=i, label=label, key=key, max_marks=max_marks, qtype=qtype,
                correct_option=_option(get("correct_option")),
                model_answer=get("model_answer"), guidance=get("guidance"),
                choice_group=get("choice_group") or None,
            ))
        except ValueError as exc:  # pydantic ValidationError is a ValueError
            reason = exc.errors()[0]["msg"] if hasattr(exc, "errors") else str(exc)
            problems.append(f"row {i} ({label}): {reason.removeprefix('Value error, ')}")

    seen: dict[str, int] = {}
    for item in items:
        if item.key in seen:
            problems.append(f"row {item.row} ({item.label}): same question as row {seen[item.key]}")
        seen.setdefault(item.key, item.row)
    groups: dict[str, set[float]] = {}
    for item in items:
        if item.choice_group:
            groups.setdefault(item.choice_group, set()).add(item.max_marks)
    for group, marks in groups.items():
        if len(marks) > 1:
            problems.append(f"choice group '{group}': alternatives must have the same max marks")
    if not items and not problems:
        problems.append("no questions found below the header row")
    if problems:
        raise SchemeError(problems)
    return MarkingScheme(subject=subject, exam=exam, items=items)
