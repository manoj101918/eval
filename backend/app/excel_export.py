"""Write approved marks to Excel: one sheet per exam, one row per student.

Only runs when every bundle of the exam is submitted and every script approved (hard rule:
no mark reaches Excel without teacher approval). Columns: Roll No | each question | Total.
"""

import os
import re
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from backend.app.marks import counted_ids, total
from backend.app.models import Exam
from backend.app.status import bundle_status
from backend.grading.scheme import MarkingScheme

_SHEET_FORBIDDEN = re.compile(r"[\[\]:*?/\\]")


class ExportError(Exception):
    def __init__(self, problems: list[str]) -> None:
        super().__init__("; ".join(problems))
        self.problems = problems


def sheet_name(exam: Exam) -> str:
    name = _SHEET_FORBIDDEN.sub("-", f"{exam.subject} {exam.class_section} {exam.name}")
    return name.strip()[:31] or f"Exam {exam.id}"


def export_problems(exam: Exam) -> list[str]:
    problems = []
    if not exam.bundles:
        problems.append("The exam has no bundles.")
    for bundle in exam.bundles:
        if bundle_status(bundle) != "submitted":
            problems.append(f"Bundle {bundle.code} is not submitted yet.")
        for script in bundle.scripts:
            if script.status != "approved":
                problems.append(f"A script in bundle {bundle.code} is not approved.")
    seen: dict[str, str] = {}
    for bundle in exam.bundles:
        for script in bundle.scripts:
            roll = (script.roll_number or "").strip().upper()
            if not roll:
                problems.append(f"A script in bundle {bundle.code} has no roll number.")
            elif roll in seen:
                problems.append(f"Roll number {roll} appears twice "
                                f"(bundles {seen[roll]} and {bundle.code}).")
            else:
                seen[roll] = bundle.code
    return list(dict.fromkeys(problems))  # keep order, drop repeats


def build_sheet(ws, exam: Exam) -> None:
    scheme = MarkingScheme.model_validate_json(exam.scheme_json)
    labels = [item.label for item in scheme.items]
    ws.append(["Roll No", *labels, "Total"])
    ws.append(["Max marks", *[item.max_marks for item in scheme.items], exam.max_marks])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for cell in ws[2]:
        cell.font = Font(italic=True)
    rows = []
    for bundle in exam.bundles:
        for script in bundle.scripts:
            by_label = {q.question: q for q in script.questions}
            counted = counted_ids(script.questions)
            values = []
            for label in labels:
                q = by_label.get(label)
                values.append(q.final_marks if q and q.id in counted else None)
            rows.append([script.roll_number.strip().upper(), *values, total(script.questions)])
    for row in sorted(rows, key=lambda r: r[0]):
        ws.append(row)
    ws.freeze_panes = "B3"
    ws.column_dimensions["A"].width = 16


def _save_atomic(wb: Workbook, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=path.parent)
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, path)  # never leave a half-written workbook
    except PermissionError:
        Path(tmp).unlink(missing_ok=True)
        raise ExportError([f"Cannot write {path.name}: close it in Excel and export again."]) \
            from None
    except BaseException:
        Path(tmp).unlink(missing_ok=True)
        raise


def export_exam(exam: Exam, *, export_dir: Path, master_path: Path | None) -> Path:
    """Write the exam's sheet and return the file path. Raises ExportError when not allowed."""
    problems = export_problems(exam)
    if problems:
        raise ExportError(problems)
    name = sheet_name(exam)
    if master_path is not None:
        if master_path.exists():
            wb = load_workbook(master_path)
        else:
            wb = Workbook()
            wb.remove(wb.active)  # drop the default empty sheet
        if name in wb.sheetnames:
            wb.remove(wb[name])  # replace this exam's previous sheet
        build_sheet(wb.create_sheet(name), exam)
        _save_atomic(wb, master_path)
        return master_path
    wb = Workbook()
    ws = wb.active
    ws.title = name
    build_sheet(ws, exam)
    slug = re.sub(r"[^A-Za-z0-9]+", "_", name).strip("_")
    path = export_dir / f"exam_{exam.id}_{slug}.xlsx"
    _save_atomic(wb, path)
    return path
